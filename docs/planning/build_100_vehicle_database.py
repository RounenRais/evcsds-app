from __future__ import annotations

import json
import re
import sys
import time
import unicodedata
from collections import Counter
from concurrent.futures import ThreadPoolExecutor, as_completed
from copy import copy
from datetime import date
from pathlib import Path
from typing import Any
from urllib.parse import urljoin

import requests
from bs4 import BeautifulSoup, Tag
from openpyxl import load_workbook
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side

from create_vehicle_database_template import FIELDS


BASE_DIR = Path(__file__).resolve().parent
TEMPLATE_PATH = BASE_DIR / "ELEKTRIKLI_ARAC_VERITABANI_SABLONU_V3.xlsx"
OUTPUT_XLSX = BASE_DIR / "ELEKTRIKLI_ARAC_VERITABANI_100_ARAC.xlsx"
OUTPUT_JSON = BASE_DIR / "ELEKTRIKLI_ARAC_VERITABANI_100_ARAC.json"

LIST_URL = "https://ev-database.org/"
SOURCE_NAME = "EV Database — European Electric Vehicle Database"
ACCESS_DATE = date(2026, 7, 24)
USER_AGENT = (
    "Mozilla/5.0 (Windows NT 10.0; Win64; x64) "
    "AppleWebKit/537.36 Chrome/126 Safari/537.36"
)

# Türkiye'de bilinen/satılan markalara ağırlık verilir. Kotalar, tek bir markanın
# örnek veri tabanını baskılamasını engeller. Toplam 100'e ulaşılamazsa aynı
# listedeki güncel araçlarla kontrollü biçimde tamamlanır.
BRAND_QUOTAS = {
    "TOGG": 6,
    "Tesla": 5,
    "Volkswagen": 5,
    "BMW": 5,
    "Mercedes-Benz": 5,
    "Audi": 4,
    "Hyundai": 4,
    "Kia": 5,
    "Renault": 4,
    "Peugeot": 4,
    "Citroen": 3,
    "Opel": 3,
    "Fiat": 3,
    "Volvo": 4,
    "Polestar": 3,
    "Skoda": 4,
    "CUPRA": 3,
    "Ford": 4,
    "Nissan": 3,
    "Toyota": 3,
    "Lexus": 2,
    "MG": 4,
    "BYD": 5,
    "Dacia": 2,
    "Mini": 3,
    "Porsche": 3,
    "Smart": 2,
    "Jeep": 1,
    "Alfa Romeo": 1,
    "Leapmotor": 2,
    "KGM": 1,
    "XPENG": 3,
}

BRAND_ALIASES = {
    "Mercedes Benz": "Mercedes-Benz",
    "Mercedes-Benz": "Mercedes-Benz",
    "Škoda": "Skoda",
    "Skoda": "Skoda",
    "Citroën": "Citroen",
    "Citroen": "Citroen",
    "MINI": "Mini",
    "mini": "Mini",
    "smart": "Smart",
    "Alfa-Romeo": "Alfa Romeo",
}

PORT_LOCATION_MAP = {
    "right side - front": "Sağ ön",
    "right side - rear": "Sağ arka",
    "left side - front": "Sol ön",
    "left side - rear": "Sol arka",
    "front side - middle": "Ön orta",
    "rear side - middle": "Arka orta",
    "front - middle": "Ön orta",
    "rear - middle": "Arka orta",
}


def clean_text(value: Any) -> str:
    return re.sub(r"\s+", " ", str(value or "")).strip()


def normalized(value: Any) -> str:
    text = unicodedata.normalize("NFKD", clean_text(value))
    return "".join(ch for ch in text if not unicodedata.combining(ch)).lower()


def number(value: Any) -> float | None:
    match = re.search(r"-?\d+(?:[.,]\d+)?", clean_text(value).replace(",", ""))
    return float(match.group(0)) if match else None


def integer(value: Any) -> int | None:
    parsed = number(value)
    return int(parsed) if parsed is not None else None


def first_year(value: Any) -> int | None:
    match = re.search(r"\b(20\d{2})\b", clean_text(value))
    return int(match.group(1)) if match else None


def table_rows(container: Tag | None) -> list[tuple[str, str]]:
    rows: list[tuple[str, str]] = []
    if not container:
        return rows
    for tr in container.find_all("tr"):
        cells = [clean_text(cell.get_text(" ", strip=True)) for cell in tr.find_all("td")]
        if len(cells) >= 2:
            rows.append((cells[0], cells[1]))
    return rows


def value_from_rows(rows: list[tuple[str, str]], label: str) -> str | None:
    wanted = normalized(label)
    for key, value in rows:
        if normalized(key).startswith(wanted):
            return value
    return None


def subsection_nodes(section: Tag | None, heading: str) -> list[Tag]:
    if section is None:
        return []
    wanted = normalized(heading)
    title = next(
        (
            h
            for h in section.find_all("h3", recursive=False)
            if normalized(h.get_text(" ", strip=True)).startswith(wanted)
        ),
        None,
    )
    if title is None:
        return []
    nodes: list[Tag] = []
    for sibling in title.find_next_siblings():
        if sibling.name == "h3":
            break
        nodes.append(sibling)
    return nodes


def subsection_rows(section: Tag | None, heading: str) -> list[tuple[str, str]]:
    rows: list[tuple[str, str]] = []
    for node in subsection_nodes(section, heading):
        rows.extend(table_rows(node))
    return rows


def map_yes_no(value: Any) -> str:
    text = normalized(value)
    if text == "yes":
        return "Evet"
    if text == "no":
        return "Hayır"
    return "Bilinmiyor"


def map_combined_support(values: list[str | None]) -> str:
    normalized_values = {normalized(value) for value in values if value}
    if "yes" in normalized_values:
        return "Evet"
    if normalized_values and normalized_values <= {"no", "-", "no data"}:
        return "Hayır"
    return "Bilinmiyor"


def map_port_location(value: Any) -> str:
    return PORT_LOCATION_MAP.get(normalized(value), "Bilinmiyor")


def normalize_brand(value: str) -> str:
    clean = clean_text(value)
    return BRAND_ALIASES.get(clean, clean)


def fetch(url: str, attempts: int = 3) -> str:
    last_error: Exception | None = None
    for attempt in range(attempts):
        try:
            response = requests.get(
                url,
                timeout=45,
                headers={"User-Agent": USER_AGENT, "Accept-Language": "en-GB,en;q=0.9"},
            )
            response.raise_for_status()
            return response.text
        except Exception as exc:  # network errors are retried and reported later
            last_error = exc
            time.sleep(1.0 + attempt)
    raise RuntimeError(f"Kaynak alınamadı: {url}: {last_error}")


def collect_candidates() -> list[dict[str, Any]]:
    soup = BeautifulSoup(fetch(LIST_URL), "html.parser")
    candidates: list[dict[str, Any]] = []
    seen_urls: set[str] = set()

    for item in soup.select("div.list-item"):
        availability = item.select_one(".availability")
        title_link = item.select_one("a.title")
        if availability is None or title_link is None:
            continue
        if "current" not in (availability.get("class") or []):
            continue

        href = clean_text(title_link.get("href"))
        if not href.startswith("/car/"):
            continue
        url = urljoin(LIST_URL, href)
        if url in seen_urls:
            continue
        seen_urls.add(url)

        brand_node = title_link.find("span", recursive=False)
        model_node = title_link.select_one("span.model")
        brand = normalize_brand(
            brand_node.get_text(" ", strip=True) if brand_node else ""
        )
        model_text = clean_text(
            model_node.get_text(" ", strip=True)
            if model_node
            else title_link.get_text(" ", strip=True).removeprefix(brand)
        )
        year_node = item.select_one(".year_from")
        year_from = integer(year_node.get_text(strip=True)) if year_node else None
        candidates.append(
            {
                "brand": brand,
                "model_text": model_text,
                "url": url,
                "year_from": year_from,
                "availability": clean_text(availability.get_text(" ", strip=True)),
            }
        )

    selected: list[dict[str, Any]] = []
    counts: Counter[str] = Counter()
    # Marka sırası bilinçlidir: Togg gibi proje sunumunda özellikle görülmesi
    # istenen markalar, kaynak sitesindeki genel sıralamadan etkilenmeden önce
    # kendi kotasını doldurur.
    for wanted_brand, quota in BRAND_QUOTAS.items():
        for candidate in candidates:
            if candidate["brand"] != wanted_brand:
                continue
            selected.append(candidate)
            counts[wanted_brand] += 1
            if counts[wanted_brand] >= quota or len(selected) == 100:
                break
        if len(selected) == 100:
            return selected

    # Kotalardan biri güncel listede dolmazsa, tercih edilen markalardaki diğer
    # güncel varyantlarla 100'e tamamlanır; marka başına üst sınır 7'dir.
    selected_urls = {item["url"] for item in selected}
    for candidate in candidates:
        brand = candidate["brand"]
        if brand not in BRAND_QUOTAS or candidate["url"] in selected_urls:
            continue
        if counts[brand] >= 7:
            continue
        selected.append(candidate)
        selected_urls.add(candidate["url"])
        counts[brand] += 1
        if len(selected) == 100:
            return selected

    raise RuntimeError(f"100 güncel araç seçilemedi; yalnızca {len(selected)} bulundu.")


def parse_home_charge_table(soup: BeautifulSoup) -> dict[str, Any]:
    section = soup.find(id="charge-table")
    if section is None:
        return {}
    tables = section.find_all("table")
    if len(tables) < 2:
        return {}

    parsed_rows: list[dict[str, Any]] = []
    for tr in tables[1].find_all("tr"):
        cells = [
            clean_text(cell.get_text(" ", strip=True))
            for cell in tr.find_all(["th", "td"])
        ]
        if len(cells) < 3 or normalized(cells[0]) == "charging point":
            continue
        phase_match = re.search(r"([13])-phase", cells[0], flags=re.I)
        current_match = re.search(r"([13])x(\d+(?:\.\d+)?)A", cells[1], flags=re.I)
        parsed_rows.append(
            {
                "label": cells[0],
                "phase_label": int(phase_match.group(1)) if phase_match else None,
                "actual_phase": int(current_match.group(1)) if current_match else None,
                "current_a": (
                    float(current_match.group(2)) if current_match else None
                ),
                "power_kw": number(cells[2]),
            }
        )

    single_rows = [
        row
        for row in parsed_rows
        if row["phase_label"] == 1 and row["actual_phase"] == 1
    ]
    three_rows = [
        row
        for row in parsed_rows
        if row["phase_label"] == 3 and row["actual_phase"] == 3
    ]
    return {"single_rows": single_rows, "three_rows": three_rows}


def best_phase_row(
    rows: list[dict[str, Any]], standard_ac_power: float | None
) -> dict[str, Any] | None:
    if not rows:
        return None
    if standard_ac_power is not None:
        standard_rows = [
            row
            for row in rows
            if row["power_kw"] is not None
            and row["power_kw"] <= standard_ac_power + 0.05
        ]
        if standard_rows:
            return max(standard_rows, key=lambda row: row["power_kw"] or -1)
    return max(rows, key=lambda row: row["power_kw"] or -1)


def parse_vehicle(candidate: dict[str, Any]) -> dict[str, Any]:
    url = candidate["url"]
    soup = BeautifulSoup(fetch(url), "html.parser")
    title = clean_text(soup.h1.get_text(" ", strip=True) if soup.h1 else "")
    brand = candidate["brand"]
    exact_variant = clean_text(title.removeprefix(brand))

    battery_rows = table_rows(soup.find(id="battery"))
    charge_section = soup.find(id="charging")
    home_rows = subsection_rows(charge_section, "Home / Destination")
    rapid_rows = subsection_rows(charge_section, "Fast Charging")
    if not rapid_rows:
        rapid_rows = subsection_rows(charge_section, "Rapid Charging")

    efficiency = soup.find(id="efficiency")
    wltp_tel_rows = subsection_rows(efficiency, "WLTP Ratings (TEL)")
    wltp_teh_rows = subsection_rows(efficiency, "WLTP Ratings (TEH)")
    wltp_rows = wltp_tel_rows or wltp_teh_rows
    if not wltp_rows:
        wltp_rows = subsection_rows(efficiency, "WLTP Ratings")

    v2x = soup.find(id="v2x")
    v2l_rows = subsection_rows(v2x, "Vehicle-to-Load")
    v2h_rows = subsection_rows(v2x, "Vehicle-to-Home")
    v2g_rows = subsection_rows(v2x, "Vehicle-to-Grid")

    ac_power = number(value_from_rows(home_rows, "Charge Power"))
    charge_matrix = parse_home_charge_table(soup)
    single = best_phase_row(charge_matrix.get("single_rows", []), ac_power)
    three = best_phase_row(charge_matrix.get("three_rows", []), ac_power)
    single_supported = single is not None
    three_supported = three is not None
    onboard_optional = "optional" in normalized(
        soup.find(id="charge-table").get_text(" ", strip=True)
        if soup.find(id="charge-table")
        else ""
    )

    dc_power = number(value_from_rows(rapid_rows, "Charge Power (max)"))
    dc_time = number(value_from_rows(rapid_rows, "Charge Time"))
    ac_port = clean_text(value_from_rows(home_rows, "Charge Port"))
    dc_port = clean_text(value_from_rows(rapid_rows, "Charge Port"))
    port_location = value_from_rows(home_rows, "Port Location")

    wltp_range = number(value_from_rows(wltp_rows, "Range"))
    rated_wh_km = number(value_from_rows(wltp_rows, "Rated Consumption"))
    wltp_consumption = rated_wh_km / 10.0 if rated_wh_km is not None else None

    v2h_values = [
        value
        for key, value in v2h_rows
        if normalized(key).endswith("supported")
    ]
    v2g_values = [
        value
        for key, value in v2g_rows
        if normalized(key).endswith("supported")
    ]

    vehicle_id_match = re.search(r"/car/(\d+)/", url)
    source_id = vehicle_id_match.group(1) if vehicle_id_match else "UNKNOWN"
    safe_slug = re.sub(r"[^A-Z0-9]+", "_", normalized(exact_variant).upper()).strip("_")
    vehicle_id = f"EU_EVDB_{source_id}_{safe_slug[:48]}"

    note_parts = [
        "Avrupa pazarındaki güncel varyant örneğidir; Türkiye donanımıyla birebir aynı olduğu varsayılmamalıdır.",
        f"Kaynak sayfasındaki tam varyant adı: {title}.",
    ]
    if "*" in clean_text(value_from_rows(battery_rows, "Useable Capacity*")):
        note_parts.append("Kullanılabilir batarya değeri kaynakta tahmini olabilir.")

    ac_note = (
        "Tek ve üç faz güç/akım değerleri kaynaktaki Home and Destination "
        "Charging tablosundan alınmıştır."
    )
    if ac_power is None:
        ac_note += " Azami AC güç açıkça bulunamadı."

    has_wltp = bool(wltp_rows and (wltp_range is not None or rated_wh_km is not None))
    battery_note = (
        "WLTP TEL değeri varsa düşük enerji tüketimli donanım sınırı kullanılmıştır; "
        "jant ve donanıma göre TEH değeri farklı olabilir."
        if has_wltp
        else "Kaynak sayfasında kullanılabilir bir WLTP değeri bulunamadı."
    )

    record = {
        "vehicle_id": vehicle_id,
        "marka": brand,
        "model": candidate["model_text"],
        "versiyon": exact_variant,
        "model_yili_baslangic": candidate.get("year_from") or first_year(title),
        "model_yili_bitis": None,
        "pazar": "Avrupa",
        "arac_tipi": "BEV",
        "uretim_durumu": "Satışta",
        "varyant_notu": " ".join(note_parts),
        "ac_konnektor_tipi": (
            "Type 2" if normalized(ac_port) == "type 2" else "Bilinmiyor"
        ),
        "ac_azami_guc_kw": ac_power,
        "ac_desteklenen_fazlar": (
            "Tek ve üç faz"
            if single_supported and three_supported
            else "Yalnızca tek faz"
            if single_supported
            else "Bilinmiyor"
        ),
        "ac_tek_faz_destegi": "Evet" if single_supported else "Bilinmiyor",
        "ac_tek_faz_azami_guc_kw": single["power_kw"] if single else None,
        "ac_tek_faz_azami_akim_a": single["current_a"] if single else None,
        "ac_uc_faz_destegi": "Evet" if three_supported else "Hayır",
        "ac_uc_faz_azami_guc_kw": three["power_kw"] if three else None,
        "ac_uc_faz_faz_basina_akim_a": three["current_a"] if three else None,
        "onboard_charger_opsiyonel_mi": "Evet" if onboard_optional else "Bilinmiyor",
        "ac_sarj_notu": ac_note,
        "batarya_brut_kwh": number(value_from_rows(battery_rows, "Nominal Capacity")),
        "batarya_kullanilabilir_kwh": number(
            value_from_rows(battery_rows, "Useable Capacity")
        ),
        "wltp_tuketim_kwh_100km": wltp_consumption,
        "wltp_menzil_km": wltp_range,
        "test_standardi": "WLTP" if has_wltp else "Bilinmiyor",
        "batarya_tuketim_notu": battery_note,
        "dc_konnektor_tipi": (
            "CCS Combo 2" if normalized(dc_port) == "ccs" else "Bilinmiyor"
        ),
        "dc_azami_guc_kw": dc_power,
        "dc_10_80_sure_dakika": dc_time,
        "zamanlanmis_sarj_destegi": "Bilinmiyor",
        "v2l_destegi": map_yes_no(value_from_rows(v2l_rows, "V2L Supported")),
        "v2h_destegi": map_combined_support(v2h_values),
        "v2g_destegi": map_combined_support(v2g_values),
        "sarj_portu_konumu": map_port_location(port_location),
        "kaynak_tipi": "Güvenilir ikincil kaynak",
        "kaynak_basligi": f"{SOURCE_NAME}: {title}",
        "kaynak_url": url,
        "erisim_tarihi": ACCESS_DATE,
        "dogrulama_durumu": "İkinci kaynak bekleniyor",
        "veri_guven_duzeyi": "Orta",
        "kaynak_notu": (
            "AC faz/akım tablosu ve teknik değerler EV Database sayfasından "
            "derlenmiştir. Üretim amaçlı kullanımdan önce ilgili üreticinin "
            "Türkiye teknik föyüyle doğrulanmalıdır."
        ),
    }
    return record


def fill_workbook(records: list[dict[str, Any]]) -> None:
    workbook = load_workbook(TEMPLATE_PATH)
    sheet = workbook.worksheets[1]
    guide = workbook.worksheets[2]

    field_names = [field[1] for field in FIELDS]
    start_row = 3
    end_row = start_row + len(records) - 1
    thin = Side(style="thin", color="D9E2F3")

    for row in range(start_row, max(sheet.max_row, end_row) + 1):
        for column in range(1, len(FIELDS) + 1):
            sheet.cell(row, column).value = None

    for row_offset, record in enumerate(records):
        row = start_row + row_offset
        fill = "E2F0D9" if row_offset % 2 == 0 else "F7FBF4"
        for column, field_name in enumerate(field_names, start=1):
            cell = sheet.cell(row, column)
            cell.value = record.get(field_name)
            cell.fill = PatternFill("solid", fgColor=fill)
            cell.border = Border(left=thin, right=thin, top=thin, bottom=thin)
            cell.alignment = Alignment(
                vertical="top",
                wrap_text=field_name
                in {
                    "varyant_notu",
                    "ac_sarj_notu",
                    "batarya_tuketim_notu",
                    "kaynak_basligi",
                    "kaynak_url",
                    "kaynak_notu",
                },
            )
            if field_name == "kaynak_url" and record.get(field_name):
                cell.hyperlink = record[field_name]
                cell.font = Font(color="0563C1", underline="single")
            elif row == start_row:
                cell.font = copy(sheet.cell(start_row + 1, column).font)
        sheet.row_dimensions[row].height = 42

    sheet.auto_filter.ref = f"A2:AP{end_row}"
    sheet.freeze_panes = "E3"
    sheet.sheet_view.showGridLines = False

    # Kullanım sayfasına kısa, sunuma uygun veri seti özeti eklenir.
    usage = workbook.worksheets[0]
    usage["A27"] = "100 araçlık örnek veri seti"
    usage["A27"].font = Font(bold=True, color="FFFFFF")
    usage["A27"].fill = PatternFill("solid", fgColor="548235")
    usage["B27"] = (
        "100 güncel Avrupa BEV varyantı; Togg ve Türkiye'de bilinen markalara "
        "ağırlık verilmiştir."
    )
    usage["A28"] = "Kaynak politikası"
    usage["A28"].font = Font(bold=True)
    usage["B28"] = (
        "Her kayıt EV Database bağlantısıyla izlenebilir. Kayıtlar üretici "
        "Türkiye teknik föyüyle ikinci doğrulama beklediği için Orta güvenlidir."
    )
    usage["A29"] = "Kritik veri kapsamı"
    usage["A29"].font = Font(bold=True)
    usage["B29"] = (
        "Araç kimliği, AC güç, tek/üç faz desteği, faz akımı, batarya ve "
        "bulunabildiği ölçüde WLTP/DC değerleri."
    )
    usage["A30"] = "Üretim uyarısı"
    usage["A30"].font = Font(bold=True, color="9C0006")
    usage["B30"] = (
        "Bu dosya danışman sunumu ve veri yapısı pilotudur; nihai karar destek "
        "kullanımından önce Türkiye varyantları üretici kaynağıyla doğrulanmalıdır."
    )
    for row in range(27, 31):
        usage[f"B{row}"].alignment = Alignment(wrap_text=True, vertical="top")
        usage.row_dimensions[row].height = 38

    # Alan rehberindeki örnek sütun başlığı, tek örnek yerine 100 kayıtlık dosyayı
    # doğru tanımlar; mevcut ayrıntılı alan açıklamaları korunur.
    if guide.max_column >= 10:
        guide.cell(1, 10).value = "İlk Araç Kaydı Örneği"

    workbook.save(OUTPUT_XLSX)


def save_json(records: list[dict[str, Any]]) -> None:
    serializable = []
    for record in records:
        converted = {
            key: value.isoformat() if isinstance(value, date) else value
            for key, value in record.items()
        }
        serializable.append(converted)
    OUTPUT_JSON.write_text(
        json.dumps(serializable, ensure_ascii=False, indent=2),
        encoding="utf-8",
    )


def postprocess_standard_ac(records: list[dict[str, Any]]) -> None:
    """Keep phase-specific values consistent with the standard AC power.

    Some source pages display standard and optional/country-specific on-board
    chargers in one matrix. The record's main Charge Power value represents the
    standard configuration, so a higher matrix value must not silently replace
    it.
    """
    for record in records:
        ac_power = record.get("ac_azami_guc_kw")
        three_power = record.get("ac_uc_faz_azami_guc_kw")
        if not isinstance(ac_power, (int, float)) or not isinstance(
            three_power, (int, float)
        ):
            continue
        if three_power <= ac_power + 0.05:
            continue

        record["onboard_charger_opsiyonel_mi"] = "Evet"
        if ac_power >= 10.5:
            record["ac_uc_faz_destegi"] = "Evet"
            record["ac_uc_faz_azami_guc_kw"] = float(ac_power)
            record["ac_uc_faz_faz_basina_akim_a"] = (
                32.0 if ac_power >= 20.0 else 16.0
            )
            record["ac_desteklenen_fazlar"] = "Tek ve üç faz"
        else:
            record["ac_uc_faz_destegi"] = "Hayır"
            record["ac_uc_faz_azami_guc_kw"] = None
            record["ac_uc_faz_faz_basina_akim_a"] = None
            record["ac_desteklenen_fazlar"] = "Yalnızca tek faz"

        record["ac_sarj_notu"] = (
            clean_text(record.get("ac_sarj_notu"))
            + " Kaynak matrisinde daha yüksek opsiyonel veya ülkeye bağlı bir "
            "şarj cihazı da gösterildiğinden standart değer ana AC güçle "
            "sınırlandırılmıştır."
        ).strip()


def validate(records: list[dict[str, Any]]) -> dict[str, Any]:
    ids = [record["vehicle_id"] for record in records]
    required = ["marka", "model", "versiyon", "ac_azami_guc_kw", "kaynak_url"]
    missing = {
        field: sum(record.get(field) in (None, "") for record in records)
        for field in required
    }
    return {
        "record_count": len(records),
        "unique_id_count": len(set(ids)),
        "brand_count": len({record["marka"] for record in records}),
        "brand_distribution": dict(sorted(Counter(r["marka"] for r in records).items())),
        "missing_required": missing,
    }


def main() -> None:
    if "--reuse-existing" in sys.argv:
        raw_records = json.loads(OUTPUT_JSON.read_text(encoding="utf-8"))
        records = [
            {
                key: (
                    date.fromisoformat(value)
                    if key == "erisim_tarihi" and isinstance(value, str)
                    else value
                )
                for key, value in record.items()
            }
            for record in raw_records
        ]
        postprocess_standard_ac(records)
        report = validate(records)
        if report["record_count"] != 100 or report["unique_id_count"] != 100:
            raise RuntimeError(f"Veri bütünlüğü hatası: {report}")
        fill_workbook(records)
        save_json(records)
        print(json.dumps(report, ensure_ascii=False, indent=2))
        print(f"Excel: {OUTPUT_XLSX}")
        print(f"JSON: {OUTPUT_JSON}")
        return

    candidates = collect_candidates()
    records_by_url: dict[str, dict[str, Any]] = {}
    errors: list[str] = []

    with ThreadPoolExecutor(max_workers=4) as executor:
        futures = {
            executor.submit(parse_vehicle, candidate): candidate
            for candidate in candidates
        }
        for future in as_completed(futures):
            candidate = futures[future]
            try:
                records_by_url[candidate["url"]] = future.result()
            except Exception as exc:
                errors.append(f"{candidate['url']}: {exc}")

    if errors:
        raise RuntimeError(
            f"{len(errors)} araç işlenemedi:\n" + "\n".join(errors[:10])
        )

    records = [records_by_url[candidate["url"]] for candidate in candidates]
    postprocess_standard_ac(records)
    report = validate(records)
    if report["record_count"] != 100 or report["unique_id_count"] != 100:
        raise RuntimeError(f"Veri bütünlüğü hatası: {report}")
    if report["missing_required"]["ac_azami_guc_kw"]:
        raise RuntimeError(f"Kritik AC güç değeri eksik: {report}")

    fill_workbook(records)
    save_json(records)
    print(json.dumps(report, ensure_ascii=False, indent=2))
    print(f"Excel: {OUTPUT_XLSX}")
    print(f"JSON: {OUTPUT_JSON}")


if __name__ == "__main__":
    main()
